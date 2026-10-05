"""Work Achievements Logger desktop application."""

from __future__ import annotations

import calendar
import json
import sqlite3
import sys
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from docx import Document
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Font

APP_DIR = Path.home() / "Documents" / "Work Achievements"
DB_PATH = APP_DIR / "work_achievements.db"
XLSX_PATH = APP_DIR / "work_achievements.xlsx"


@dataclass
class Entry:
    id: int
    completed_at: datetime
    title: str
    description: str


def parse_timestamp(value: str) -> datetime:
    return datetime.fromisoformat(value)


def week_start_for(value: date) -> date:
    return value - timedelta(days=value.weekday())


def period_bounds(view: str, selected: date) -> tuple[date, date]:
    if view == "Daily":
        return selected, selected
    if view == "Weekly":
        start = week_start_for(selected)
        return start, start + timedelta(days=6)
    last_day = calendar.monthrange(selected.year, selected.month)[1]
    return selected.replace(day=1), selected.replace(day=last_day)


def period_label(view: str, selected: date) -> str:
    start, end = period_bounds(view, selected)
    if view == "Daily":
        return start.strftime("%A, %B %-d, %Y") if sys.platform != "win32" else start.strftime("%A, %B %#d, %Y")
    if view == "Weekly":
        return f"Week of {start:%B} {start.day}, {start.year} through {end:%B} {end.day}, {end.year}"
    return selected.strftime("%B %Y")


class Repository:
    def __init__(self, path: Path = DB_PATH):
        APP_DIR.mkdir(parents=True, exist_ok=True)
        self.path = path
        self.connection = sqlite3.connect(path)
        self.connection.row_factory = sqlite3.Row
        self.connection.execute(
            """CREATE TABLE IF NOT EXISTS entries (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                completed_at TEXT NOT NULL,
                title TEXT NOT NULL,
                description TEXT NOT NULL
            )"""
        )
        self.connection.commit()

    def add(self, title: str, description: str) -> Entry:
        completed_at = datetime.now().replace(microsecond=0)
        cursor = self.connection.execute(
            "INSERT INTO entries(completed_at, title, description) VALUES (?, ?, ?)",
            (completed_at.isoformat(), title, description),
        )
        self.connection.commit()
        return Entry(cursor.lastrowid, completed_at, title, description)

    def update(self, entry_id: int, title: str, description: str) -> None:
        self.connection.execute(
            "UPDATE entries SET title = ?, description = ? WHERE id = ?",
            (title, description, entry_id),
        )
        self.connection.commit()

    def delete(self, entry_id: int) -> None:
        self.connection.execute("DELETE FROM entries WHERE id = ?", (entry_id,))
        self.connection.commit()

    def all(self) -> list[Entry]:
        rows = self.connection.execute(
            "SELECT id, completed_at, title, description FROM entries ORDER BY completed_at DESC"
        ).fetchall()
        return [Entry(row["id"], parse_timestamp(row["completed_at"]), row["title"], row["description"]) for row in rows]

    def close(self) -> None:
        self.connection.close()


def entries_for(entries: list[Entry], view: str, selected: date) -> list[Entry]:
    start, end = period_bounds(view, selected)
    return [entry for entry in entries if start <= entry.completed_at.date() <= end]


def export_excel(entries: list[Entry], path: Path = XLSX_PATH) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Entries"
    headers = ["Completed date", "Completed time", "Title", "Description", "Week", "Month"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for entry in sorted(entries, key=lambda item: item.completed_at):
        sheet.append(
            [
                entry.completed_at.date().isoformat(),
                entry.completed_at.strftime("%H:%M:%S"),
                entry.title,
                entry.description,
                week_start_for(entry.completed_at.date()).isoformat(),
                entry.completed_at.strftime("%Y-%m"),
            ]
        )
    widths = {"A": 18, "B": 15, "C": 30, "D": 70, "E": 18, "F": 12}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    workbook.save(path)


def export_word(entries: list[Entry], view: str, selected: date, path: Path) -> None:
    document = Document()
    document.add_heading("Work Achievements", level=0)
    document.add_paragraph(period_label(view, selected))
    if not entries:
        document.add_paragraph("No entries were recorded for this period.")
    else:
        table = document.add_table(rows=1, cols=3)
        table.style = "Light Shading Accent 1"
        headings = ["Completed", "Title", "Description"]
        for cell, heading in zip(table.rows[0].cells, headings):
            cell.text = heading
        for entry in sorted(entries, key=lambda item: item.completed_at):
            cells = table.add_row().cells
            cells[0].text = entry.completed_at.strftime("%Y-%m-%d %H:%M")
            cells[1].text = entry.title
            cells[2].text = entry.description
    document.save(path)


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Work Achievements Logger")
        self.geometry("980x650")
        self.minsize(800, 500)
        self.repository = Repository()
        self.selected_id: int | None = None
        self.view = tk.StringVar(value="Daily")
        self.selected_date = tk.StringVar(value=date.today().isoformat())
        self.title_value = tk.StringVar()
        self._build_ui()
        self.refresh()
        self.protocol("WM_DELETE_WINDOW", self.on_close)

    def _build_ui(self):
        form = ttk.LabelFrame(self, text="Record completed work", padding=10)
        form.pack(fill="x", padx=10, pady=(10, 5))
        ttk.Label(form, text="Title:").grid(row=0, column=0, sticky="w", padx=(0, 6), pady=4)
        ttk.Entry(form, textvariable=self.title_value, width=50).grid(row=0, column=1, sticky="ew", pady=4)
        ttk.Label(form, text="Description:").grid(row=1, column=0, sticky="nw", padx=(0, 6), pady=4)
        self.description = tk.Text(form, height=4, width=60, wrap="word")
        self.description.grid(row=1, column=1, sticky="ew", pady=4)
        self.save_button = ttk.Button(form, text="Add entry", command=self.save_entry)
        self.save_button.grid(row=2, column=1, sticky="w", pady=(8, 0))
        ttk.Button(form, text="Clear", command=self.clear_form).grid(row=2, column=1, sticky="w", padx=(95, 0), pady=(8, 0))
        form.columnconfigure(1, weight=1)

        controls = ttk.Frame(self, padding=(10, 5))
        controls.pack(fill="x")
        ttk.Label(controls, text="Organize by:").pack(side="left")
        selector = ttk.Combobox(controls, textvariable=self.view, values=("Daily", "Weekly", "Monthly"), state="readonly", width=12)
        selector.pack(side="left", padx=6)
        selector.bind("<<ComboboxSelected>>", lambda _event: self.refresh())
        ttk.Label(controls, text="Date (YYYY-MM-DD):").pack(side="left", padx=(12, 4))
        date_entry = ttk.Entry(controls, textvariable=self.selected_date, width=13)
        date_entry.pack(side="left")
        date_entry.bind("<Return>", lambda _event: self.refresh())
        ttk.Button(controls, text="Refresh", command=self.refresh).pack(side="left", padx=6)
        ttk.Button(controls, text="Export current view to Word", command=self.export_current_word).pack(side="right")

        self.period_text = ttk.Label(self, padding=(10, 0))
        self.period_text.pack(anchor="w")
        table_frame = ttk.Frame(self, padding=10)
        table_frame.pack(fill="both", expand=True)
        columns = ("completed", "title", "description")
        self.table = ttk.Treeview(table_frame, columns=columns, show="headings", selectmode="browse")
        self.table.heading("completed", text="Completed")
        self.table.heading("title", text="Title")
        self.table.heading("description", text="Description")
        self.table.column("completed", width=150, anchor="w")
        self.table.column("title", width=250, anchor="w")
        self.table.column("description", width=500, anchor="w")
        self.table.pack(side="left", fill="both", expand=True)
        scrollbar = ttk.Scrollbar(table_frame, orient="vertical", command=self.table.yview)
        scrollbar.pack(side="right", fill="y")
        self.table.configure(yscrollcommand=scrollbar.set)
        self.table.bind("<<TreeviewSelect>>", self.select_entry)

        actions = ttk.Frame(self, padding=(10, 0, 10, 10))
        actions.pack(fill="x")
        ttk.Button(actions, text="Edit selected", command=self.edit_entry).pack(side="left")
        ttk.Button(actions, text="Delete selected", command=self.delete_entry).pack(side="left", padx=6)
        ttk.Label(actions, text=f"Files are saved in: {APP_DIR}").pack(side="right")

    def selected_period(self) -> date:
        try:
            return date.fromisoformat(self.selected_date.get().strip())
        except ValueError as error:
            raise ValueError("Enter a valid date in YYYY-MM-DD format.") from error

    def refresh(self):
        try:
            selected = self.selected_period()
        except ValueError as error:
            messagebox.showerror("Invalid date", str(error))
            return
        entries = entries_for(self.repository.all(), self.view.get(), selected)
        self.period_text.configure(text=f"{self.view.get()} view: {period_label(self.view.get(), selected)} ({len(entries)} entries)")
        for item in self.table.get_children():
            self.table.delete(item)
        for entry in entries:
            self.table.insert("", "end", iid=str(entry.id), values=(entry.completed_at.strftime("%Y-%m-%d %H:%M"), entry.title, entry.description))
        self.selected_id = None

    def select_entry(self, _event=None):
        selection = self.table.selection()
        self.selected_id = int(selection[0]) if selection else None

    def save_entry(self):
        title = self.title_value.get().strip()
        description = self.description.get("1.0", "end").strip()
        if not title or not description:
            messagebox.showwarning("Missing information", "Title and description are required.")
            return
        self.repository.add(title, description)
        self.sync_excel()
        self.clear_form()
        self.selected_date.set(date.today().isoformat())
        self.refresh()

    def edit_entry(self):
        if self.selected_id is None:
            messagebox.showinfo("Select an entry", "Select an entry before editing it.")
            return
        title = self.title_value.get().strip()
        description = self.description.get("1.0", "end").strip()
        if not title or not description:
            messagebox.showwarning("Missing information", "Title and description are required.")
            return
        self.repository.update(self.selected_id, title, description)
        self.sync_excel()
        self.clear_form()
        self.refresh()

    def delete_entry(self):
        if self.selected_id is None:
            messagebox.showinfo("Select an entry", "Select an entry before deleting it.")
            return
        if messagebox.askyesno("Delete entry", "Delete the selected entry?"):
            self.repository.delete(self.selected_id)
            self.sync_excel()
            self.clear_form()
            self.refresh()

    def clear_form(self):
        self.title_value.set("")
        self.description.delete("1.0", "end")
        self.selected_id = None
        self.save_button.configure(text="Add entry")

    def sync_excel(self):
        try:
            export_excel(self.repository.all())
        except Exception as error:
            messagebox.showerror("Excel export failed", f"The entry was saved, but Excel could not be updated:\n\n{error}")

    def export_current_word(self):
        try:
            selected = self.selected_period()
        except ValueError as error:
            messagebox.showerror("Invalid date", str(error))
            return
        path = filedialog.asksaveasfilename(
            title="Save Word report",
            initialdir=APP_DIR,
            initialfile=f"work_achievements_{self.view.get().lower()}_{selected.isoformat()}.docx",
            defaultextension=".docx",
            filetypes=(("Word document", "*.docx"), ("All files", "*.*")),
        )
        if not path:
            return
        try:
            export_word(entries_for(self.repository.all(), self.view.get(), selected), self.view.get(), selected, Path(path))
            messagebox.showinfo("Export complete", f"Word report saved to:\n{path}")
        except Exception as error:
            messagebox.showerror("Word export failed", str(error))

    def on_close(self):
        self.repository.close()
        self.destroy()


if __name__ == "__main__":
    App().mainloop()
